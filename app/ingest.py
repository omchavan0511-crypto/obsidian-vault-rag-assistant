import os
import frontmatter
from pypdf import PdfReader
from docx import Document as DocxDocument
from llama_index.core import Document, VectorStoreIndex, StorageContext
from llama_index.core.node_parser import TokenTextSplitter
from llama_index.vector_stores.chroma import ChromaVectorStore
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
import chromadb

EMBED_MODEL = HuggingFaceEmbedding(model_name="BAAI/bge-small-en-v1.5")
SKIP_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg", ".ico",
    ".mp3", ".mp4", ".mov", ".avi", ".wav",
    ".zip", ".exe", ".bin", ".db", ".sqlite",
}


def read_md(filepath: str):
    post = frontmatter.load(filepath)
    return post.content, post.get("tags", [])


def read_txt(filepath: str):
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def read_pdf(filepath: str):
    reader = PdfReader(filepath)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def read_docx(filepath: str):
    doc = DocxDocument(filepath)
    return "\n".join(p.text for p in doc.paragraphs)


def load_file(filepath: str):
    ext = os.path.splitext(filepath)[1].lower()
    if ext in SKIP_EXTENSIONS:
        return None, []

    try:
        if ext == ".md":
            return read_md(filepath)
        elif ext == ".pdf":
            return read_pdf(filepath), []
        elif ext == ".docx":
            return read_docx(filepath), []
        else:
            return read_txt(filepath), []
    except Exception:
        return None, []


def load_vault(vault_path: str) -> list[Document]:
    documents = []
    for root, _, filenames in os.walk(vault_path):
        for filename in filenames:
            filepath = os.path.join(root, filename)
            text, tags = load_file(filepath)
            if not text or not text.strip():
                continue

            title = os.path.splitext(filename)[0]
            documents.append(
                Document(
                    text=text,
                    metadata={
                        "file_name": filename,
                        "title": title,
                        "path": filepath,
                        "tags": ", ".join(tags),
                    },
                )
            )
    return documents


def build_index(vault_path: str, persist_dir: str) -> VectorStoreIndex:
    documents = load_vault(vault_path)
    if not documents:
        raise ValueError(f"No readable text files found under {vault_path}")

    chroma_client = chromadb.PersistentClient(path=persist_dir)
    collection = chroma_client.get_or_create_collection("obsidian_vault")
    vector_store = ChromaVectorStore(chroma_collection=collection)
    storage_context = StorageContext.from_defaults(vector_store=vector_store)

    index = VectorStoreIndex.from_documents(
        documents,
        storage_context=storage_context,
        embed_model=EMBED_MODEL,
        transformations=[TokenTextSplitter(chunk_size=512, chunk_overlap=50)],
    )
    return index


def load_index(persist_dir: str) -> VectorStoreIndex:
    chroma_client = chromadb.PersistentClient(path=persist_dir)
    collection = chroma_client.get_or_create_collection("obsidian_vault")
    vector_store = ChromaVectorStore(chroma_collection=collection)
    return VectorStoreIndex.from_vector_store(
        vector_store,
        embed_model=EMBED_MODEL,
    )